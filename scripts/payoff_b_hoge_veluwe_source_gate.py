#!/usr/bin/env python3
"""Materialize and audit the registered Hoge Veluwe Gate-A sources.

This script is intentionally source-readiness only. It may download and inspect
individual source files to certify immutable hashes, file identities, schemas
and year coverage. It must not join the migrant, resident, resource and cue
coordinates, calculate focal-partner mismatch, calculate cue-resource
connectivity, choose a breakpoint, or run the registered history test.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import shutil
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qsl, quote, urlencode, urljoin, urlparse, urlunparse

ROOT = Path(__file__).resolve().parents[1]
DRYAD_API = "https://datadryad.org/api/v2"
DRYAD_WEB = "https://datadryad.org"
DRYAD_PUBLIC_MIRRORS = {
    "10.5061/dryad.f1vhhmgx6": {
        "provider": "zenodo",
        "record_id": 5730499,
    }
}

YEAR_HEADER_RE = re.compile(r"(?i)(year|yr|jaar)")
FILENAME_RE = re.compile(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)', re.I)


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument(
        "--contract",
        type=Path,
        default=ROOT
        / "data"
        / "payoff_b_hoge_veluwe_network_hysteresis_contract_20260927.json",
    )
    p.add_argument(
        "--output-dir",
        type=Path,
        default=Path("outputs/hoge_veluwe_source_gate_a"),
    )
    p.add_argument(
        "--cue-receipt",
        type=Path,
        default=None,
        help="receipt from the frozen 1980-2015 Ivory Coast cue extension",
    )
    p.add_argument(
        "--timeout",
        type=int,
        default=120,
    )
    return p.parse_args()


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def sha256_path(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _session():
    import requests
    from requests.adapters import HTTPAdapter
    from urllib3.util.retry import Retry

    session = requests.Session()
    session.headers.update(
        {"User-Agent": "PAYOFF-B-Hoge-Veluwe-source-gate/1.0"}
    )
    retry = Retry(
        total=6,
        connect=4,
        read=4,
        status=6,
        backoff_factor=1.0,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET"}),
        respect_retry_after_header=True,
    )
    adapter = HTTPAdapter(
        max_retries=retry,
        pool_connections=2,
        pool_maxsize=2,
    )
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


def _api_url(href: str) -> str:
    if href.startswith("http://") or href.startswith("https://"):
        return href
    return urljoin(DRYAD_WEB, href)


def _get_json(session, url: str, timeout: int) -> dict:
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    return response.json()


def _dryad_latest_files(session, doi: str, timeout: int):
    encoded = quote(f"doi:{doi}", safe="")
    versions_url = f"{DRYAD_API}/datasets/{encoded}/versions"
    versions = _get_json(session, versions_url, timeout)
    rows = (
        versions.get("_embedded", {})
        .get("stash:versions", [])
    )
    if not rows:
        raise RuntimeError(f"Dryad DOI has no visible versions: {doi}")

    def version_key(row):
        return (
            int(row.get("versionNumber") or 0),
            str(row.get("publicationDate") or ""),
        )

    latest = max(rows, key=version_key)
    files_href = latest.get("_links", {}).get("stash:files", {}).get("href")
    if not files_href:
        raise RuntimeError(f"Dryad version has no files link: {doi}")

    files = []
    page = 1
    while True:
        separator = "&" if "?" in files_href else "?"
        url = _api_url(files_href) + f"{separator}per_page=100&page={page}"
        payload = _get_json(session, url, timeout)
        chunk = (
            payload.get("_embedded", {})
            .get("stash:files", [])
        )
        files.extend(chunk)
        total = int(payload.get("total") or len(files))
        if len(files) >= total:
            break
        if not chunk:
            raise RuntimeError(
                f"Dryad file pagination stopped early for {doi}: "
                f"{len(files)}/{total}"
            )
        page += 1
    return latest, files


def _dryad_dataset_archive_download(
    session,
    *,
    doi: str,
    exact_name: str,
    output_dir: Path,
    timeout: int,
) -> dict:
    encoded_once = quote(f"doi:{doi}", safe="")
    encoded_twice = quote(encoded_once, safe="")
    candidates = [
        f"{DRYAD_API}/datasets/{encoded_once}/download",
        f"{DRYAD_API}/datasets/{encoded_twice}/download",
        f"http://datadryad.org/api/v2/datasets/{encoded_once}/download",
        f"http://datadryad.org/api/v2/datasets/{encoded_twice}/download",
    ]
    errors = []
    for url in candidates:
        try:
            response = session.get(url, timeout=timeout, allow_redirects=True)
            if not response.ok or not response.content:
                errors.append(f"{url}: HTTP {response.status_code}")
                continue
            try:
                with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                    matches = [
                        member
                        for member in archive.namelist()
                        if Path(member).name == exact_name
                    ]
                    if len(matches) != 1:
                        errors.append(
                            f"{url}: exact file count {len(matches)} "
                            f"for {exact_name!r}"
                        )
                        continue
                    payload = archive.read(matches[0])
            except zipfile.BadZipFile:
                errors.append(f"{url}: response was not a ZIP archive")
                continue

            target = output_dir / exact_name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(payload)
            return {
                "provider": "dryad_dataset_archive",
                "doi": doi,
                "filename": exact_name,
                "download_url_used": response.url,
                "dataset_archive_sha256": sha256_bytes(response.content),
                "downloaded_size": target.stat().st_size,
                "sha256": sha256_path(target),
                "path": str(target),
            }
        except Exception as exc:
            errors.append(f"{url}: {type(exc).__name__}: {exc}")
    raise RuntimeError(
        f"Dryad dataset archive could not supply {exact_name!r}: {errors}"
    )


def _zenodo_mirror_download(
    session,
    *,
    record_id: int,
    exact_name: str,
    output_dir: Path,
    timeout: int,
) -> dict:
    api_url = f"https://zenodo.org/api/records/{record_id}"
    payload = _get_json(session, api_url, timeout)
    files = payload.get("files", [])
    matches = [row for row in files if row.get("key") == exact_name]
    if len(matches) != 1:
        raise RuntimeError(
            f"Zenodo mirror {record_id} has {len(matches)} matches for "
            f"{exact_name!r}; available={[row.get('key') for row in files]}"
        )
    meta = matches[0]
    links = meta.get("links", {})
    download_url = links.get("content") or links.get("self")
    if not download_url:
        raise RuntimeError(
            f"Zenodo mirror {record_id} has no download URL for {exact_name!r}"
        )
    response = session.get(download_url, timeout=timeout, allow_redirects=True)
    response.raise_for_status()
    if not response.content:
        raise RuntimeError(
            f"Zenodo mirror {record_id} returned empty bytes for {exact_name!r}"
        )
    target = output_dir / exact_name
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(response.content)
    return {
        "provider": "zenodo",
        "record_id": record_id,
        "filename": exact_name,
        "download_url_used": response.url,
        "downloaded_size": target.stat().st_size,
        "zenodo_checksum": meta.get("checksum"),
        "sha256": sha256_path(target),
        "path": str(target),
    }


def _dryad_download_file(
    session,
    doi: str,
    exact_name: str,
    output_dir: Path,
    timeout: int,
) -> dict:
    latest, files = _dryad_latest_files(session, doi, timeout)
    matches = [row for row in files if row.get("path") == exact_name]
    if len(matches) != 1:
        names = [row.get("path") for row in files]
        raise RuntimeError(
            f"expected one Dryad file {exact_name!r}, found {len(matches)}; "
            f"available={names}"
        )
    meta = matches[0]
    self_href = meta.get("_links", {}).get("self", {}).get("href", "")
    file_id = self_href.rstrip("/").split("/")[-1]
    download_href = (
        meta.get("_links", {})
        .get("stash:download", {})
        .get("href")
    )
    candidates = []
    if download_href:
        candidates.append(_api_url(download_href))
    if file_id:
        # Public Dryad landing pages expose file bytes through /downloads/
        # rather than the authenticated API /files/{id}/download endpoint.
        candidates.extend(
            [
                f"{DRYAD_WEB}/downloads/file_stream/{file_id}",
                f"http://datadryad.org/downloads/file_stream/{file_id}",
            ]
        )

    response = None
    errors = []
    for url in candidates:
        try:
            trial = session.get(url, timeout=timeout, allow_redirects=True)
            if trial.ok and trial.content:
                response = trial
                break
            errors.append(f"{url}: HTTP {trial.status_code}")
        except Exception as exc:
            errors.append(f"{url}: {type(exc).__name__}: {exc}")
    archive_copy = None
    mirror = None
    if response is None:
        try:
            archive_copy = _dryad_dataset_archive_download(
                session,
                doi=doi,
                exact_name=exact_name,
                output_dir=output_dir,
                timeout=timeout,
            )
        except Exception as exc:
            errors.append(
                "dryad dataset archive: "
                f"{type(exc).__name__}: {exc}"
            )

    if response is None and archive_copy is None:
        mirror_spec = DRYAD_PUBLIC_MIRRORS.get(doi)
        if mirror_spec and mirror_spec.get("provider") == "zenodo":
            try:
                mirror = _zenodo_mirror_download(
                    session,
                    record_id=int(mirror_spec["record_id"]),
                    exact_name=exact_name,
                    output_dir=output_dir,
                    timeout=timeout,
                )
            except Exception as exc:
                errors.append(
                    "zenodo mirror "
                    f"{mirror_spec.get('record_id')}: "
                    f"{type(exc).__name__}: {exc}"
                )
        if mirror is None:
            raise RuntimeError(
                f"could not download Dryad file {exact_name!r}: {errors}"
            )

    target = output_dir / exact_name
    target.parent.mkdir(parents=True, exist_ok=True)
    if response is not None:
        target.write_bytes(response.content)
    # archive_copy and mirror already wrote the exact registered file.
    digest = sha256_path(target)

    declared_digest = meta.get("digest")
    declared_type = str(meta.get("digestType") or "").lower()
    digest_match = None
    has_authoritative_sha256 = (
        bool(declared_digest)
        and declared_type in {"sha-256", "sha256"}
    )
    if mirror is not None and not has_authoritative_sha256:
        raise RuntimeError(
            f"Dryad metadata lacks authoritative SHA-256 for mirror validation "
            f"of {exact_name}: digestType={meta.get('digestType')!r}"
        )
    if has_authoritative_sha256:
        digest_match = digest.lower() == str(declared_digest).lower()
        if not digest_match:
            raise RuntimeError(
                f"Dryad SHA-256 mismatch for {exact_name}: "
                f"computed={digest} declared={declared_digest}"
            )

    return {
        "source_kind": "dryad",
        "doi": doi,
        "version_number": latest.get("versionNumber"),
        "publication_date": latest.get("publicationDate"),
        "file_id": file_id,
        "filename": exact_name,
        "declared_size": meta.get("size"),
        "downloaded_size": target.stat().st_size,
        "mime_type": meta.get("mimeType"),
        "declared_digest": declared_digest,
        "declared_digest_type": meta.get("digestType"),
        "sha256": digest,
        "declared_digest_match": digest_match,
        "download_url_used": (
            response.url
            if response is not None
            else (
                archive_copy["download_url_used"]
                if archive_copy is not None
                else mirror["download_url_used"]
            )
        ),
        "download_transport": (
            "dryad_individual"
            if response is not None
            else (
                "dryad_dataset_archive"
                if archive_copy is not None
                else "zenodo_mirror_digest_verified"
            )
        ),
        "dryad_dataset_archive": (
            None
            if archive_copy is None
            else {
                "dataset_archive_sha256": archive_copy["dataset_archive_sha256"],
                "file_sha256": archive_copy["sha256"],
            }
        ),
        "mirror": (
            None
            if mirror is None
            else {
                "provider": mirror["provider"],
                "record_id": mirror["record_id"],
                "zenodo_checksum": mirror.get("zenodo_checksum"),
                "sha256": mirror["sha256"],
            }
        ),
        "path": str(target),
    }



class _MdaLandingParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.clicks = []
        self.forms = []
        self._form = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        href = attrs.get("href")
        if href:
            self.links.append(href)
        onclick = attrs.get("onclick")
        if onclick:
            self.clicks.append(onclick)
        if tag.lower() == "form":
            self._form = {
                "method": str(attrs.get("method") or "get").lower(),
                "action": attrs.get("action") or "",
                "fields": {},
            }
        elif tag.lower() == "input" and self._form is not None:
            name = attrs.get("name")
            if name:
                self._form["fields"][name] = attrs.get("value") or ""

    def handle_endtag(self, tag):
        if tag.lower() == "form" and self._form is not None:
            self.forms.append(self._form)
            self._form = None


def _looks_html(response) -> bool:
    content_type = str(response.headers.get("Content-Type") or "").lower()
    prefix = response.content[:256].lstrip().lower()
    return (
        "text/html" in content_type
        or prefix.startswith(b"<!doctype html")
        or prefix.startswith(b"<html")
    )


def _landing_filename(html: str) -> str | None:
    plain = re.sub(r"<[^>]+>", " ", html)
    match = re.search(r"(?is)\bFile\s*:\s*['\"]([^'\"]+)['\"]", plain)
    return match.group(1).strip() if match else None


def _quoted_url_candidates(script: str) -> list[str]:
    rows = []
    for value in re.findall(r"""['"]([^'"]+)['"]""", script):
        low = value.lower()
        if (
            "download" in low
            or "getfile" in low
            or "file=" in low
            or "fid=" in low
            or low.endswith((".csv", ".tsv", ".txt", ".xlsx", ".xls", ".zip"))
        ):
            rows.append(value)
    return rows


def _append_query(url: str, fields: dict[str, str]) -> str:
    parsed = urlparse(url)
    query = list(parse_qsl(parsed.query, keep_blank_values=True))
    query.extend((str(k), str(v)) for k, v in fields.items())
    return urlunparse(parsed._replace(query=urlencode(query)))


def _resolve_mda_landing_download(
    session,
    response,
    *,
    registered_url: str,
    timeout: int,
):
    html = response.content.decode(
        response.encoding or "utf-8",
        errors="replace",
    )
    parser = _MdaLandingParser()
    parser.feed(html)
    landing_name = _landing_filename(html)

    attempts = []
    seen = set()

    def accept(trial, route: str):
        if not trial.ok or not trial.content:
            attempts.append(f"{route}: HTTP {trial.status_code}")
            return None
        if _looks_html(trial):
            attempts.append(f"{route}: HTML landing page")
            return None
        return trial

    candidates = []
    for href in parser.links:
        low = href.lower()
        if (
            "download" in low
            or "getfile" in low
            or "file=" in low
            or "fid=" in low
            or low.endswith((".csv", ".tsv", ".txt", ".xlsx", ".xls", ".zip"))
        ):
            candidates.append(urljoin(response.url, href))
    for script in parser.clicks:
        for value in _quoted_url_candidates(script):
            candidates.append(urljoin(response.url, value))

    for url in candidates:
        if url in seen or url == registered_url or url == response.url:
            continue
        seen.add(url)
        try:
            trial = session.get(url, timeout=timeout, allow_redirects=True)
            accepted = accept(trial, f"GET {url}")
            if accepted is not None:
                return accepted, landing_name, {
                    "landing_url": response.url,
                    "resolution": "html_link_or_onclick",
                    "attempts": attempts,
                }
        except Exception as exc:
            attempts.append(f"GET {url}: {type(exc).__name__}: {exc}")

    for form in parser.forms:
        action = urljoin(response.url, form["action"] or response.url)
        method = form["method"]
        fields = form["fields"]
        try:
            if method == "post":
                trial = session.post(
                    action,
                    data=fields,
                    timeout=timeout,
                    allow_redirects=True,
                )
                route = f"POST {action}"
            else:
                target = _append_query(action, fields)
                trial = session.get(
                    target,
                    timeout=timeout,
                    allow_redirects=True,
                )
                route = f"GET {target}"
            accepted = accept(trial, route)
            if accepted is not None:
                return accepted, landing_name, {
                    "landing_url": response.url,
                    "resolution": f"html_form_{method}",
                    "attempts": attempts,
                }
        except Exception as exc:
            attempts.append(
                f"{method.upper()} {action}: {type(exc).__name__}: {exc}"
            )

    probe = {
        "landing_filename": landing_name,
        "href_count": len(parser.links),
        "form_count": len(parser.forms),
        "onclick_count": len(parser.clicks),
        "hrefs": parser.links[:20],
        "forms": parser.forms[:10],
        "onclicks": parser.clicks[:10],
        "attempts": attempts[-20:],
    }
    raise RuntimeError(
        "MDA landing page did not expose a retrievable file route; "
        + json.dumps(probe, sort_keys=True)
    )


def _content_disposition_filename(headers) -> str | None:
    value = headers.get("Content-Disposition", "")
    match = FILENAME_RE.search(value)
    if not match:
        return None
    return match.group(1).strip().strip('"')


def _mda_download_file(session, url: str, output_dir: Path, timeout: int) -> dict:
    response = session.get(url, timeout=timeout, allow_redirects=True)
    response.raise_for_status()

    landing = None
    landing_name = None
    if _looks_html(response):
        response, landing_name, landing = _resolve_mda_landing_download(
            session,
            response,
            registered_url=url,
            timeout=timeout,
        )
        response.raise_for_status()

    payload = response.content
    if not payload:
        raise RuntimeError("Marine Data Archive returned an empty source file")

    name = _content_disposition_filename(response.headers) or landing_name
    content_type = response.headers.get("Content-Type", "")
    if not name:
        parsed_name = Path(urlparse(response.url).path).name
        if parsed_name and "." in parsed_name:
            name = parsed_name
        elif payload[:4] == b"PK\x03\x04":
            name = "tomotani_migrant_source.xlsx"
        elif payload[:2] == b"\x1f\x8b":
            name = "tomotani_migrant_source.gz"
        else:
            name = "tomotani_migrant_source.dat"

    target = output_dir / name
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(payload)
    return {
        "source_kind": "marine_data_archive",
        "registered_url": url,
        "download_url_used": response.url,
        "filename": name,
        "downloaded_size": target.stat().st_size,
        "content_type": content_type,
        "sha256": sha256_path(target),
        "landing_resolution": landing,
        "path": str(target),
    }


def _normalized_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _year_value(value) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    rounded = int(round(number))
    if abs(number - rounded) > 1e-6:
        return None
    if 1900 <= rounded <= 2100:
        return rounded
    return None


def _inspect_xlsx_bytes(payload: bytes) -> dict:
    import openpyxl

    workbook = openpyxl.load_workbook(
        io.BytesIO(payload),
        read_only=True,
        data_only=True,
    )
    sheets = []
    for ws in workbook.worksheets:
        rows = ws.iter_rows(values_only=True)
        first_rows = []
        for _ in range(10):
            try:
                first_rows.append(next(rows))
            except StopIteration:
                break
        header_index = None
        header = None
        for idx, row in enumerate(first_rows):
            values = [_normalized_header(v) for v in row]
            nonempty = [v for v in values if v]
            if len(nonempty) >= 2:
                header_index = idx
                header = values
                break
        if header is None:
            sheets.append(
                {
                    "sheet": ws.title,
                    "header_row": None,
                    "columns": [],
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                    "year_columns": [],
                }
            )
            continue

        year_columns = []
        for col_idx, name in enumerate(header):
            if not name or not YEAR_HEADER_RE.search(name):
                continue
            years = []
            # Include any rows already consumed after the header.
            for row in first_rows[header_index + 1 :]:
                if col_idx < len(row):
                    year = _year_value(row[col_idx])
                    if year is not None:
                        years.append(year)
            # Continue streaming the rest of the worksheet.
            for row in rows:
                if col_idx < len(row):
                    year = _year_value(row[col_idx])
                    if year is not None:
                        years.append(year)
            year_columns.append(
                {
                    "column": name,
                    "unique_year_count": len(set(years)),
                    "min_year": min(years) if years else None,
                    "max_year": max(years) if years else None,
                }
            )
            # Worksheet iterator is consumed by the first year column. There is
            # normally only one calendar-year column in these registered files.
            break

        sheets.append(
            {
                "sheet": ws.title,
                "header_row": header_index + 1,
                "columns": [name for name in header if name],
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "year_columns": year_columns,
            }
        )
    return {
        "format": "xlsx",
        "sheets": sheets,
    }


def _inspect_delimited_text(payload: bytes) -> dict:
    import csv

    text = payload.decode("utf-8-sig")
    sample = text[:8192]
    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        return {"format": "delimited_text", "columns": [], "year_columns": []}
    header = [_normalized_header(v) for v in rows[0]]
    year_columns = []
    for col_idx, name in enumerate(header):
        if not name or not YEAR_HEADER_RE.search(name):
            continue
        years = []
        for row in rows[1:]:
            if col_idx < len(row):
                year = _year_value(row[col_idx])
                if year is not None:
                    years.append(year)
        year_columns.append(
            {
                "column": name,
                "unique_year_count": len(set(years)),
                "min_year": min(years) if years else None,
                "max_year": max(years) if years else None,
            }
        )
    return {
        "format": "delimited_text",
        "columns": [name for name in header if name],
        "row_count": len(rows),
        "year_columns": year_columns,
    }


def inspect_source(path: Path) -> dict:
    payload = path.read_bytes()
    try:
        return _inspect_xlsx_bytes(payload)
    except Exception as xlsx_error:
        xlsx_error_text = f"{type(xlsx_error).__name__}: {xlsx_error}"

    try:
        if b"\x00" not in payload[:4096]:
            result = _inspect_delimited_text(payload)
            result["xlsx_open_error"] = xlsx_error_text
            return result
    except Exception as text_error:
        text_error_text = f"{type(text_error).__name__}: {text_error}"
    else:
        text_error_text = "binary_or_not_delimited"

    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            members = sorted(archive.namelist())
        return {
            "format": "zip",
            "members": members,
            "xlsx_open_error": xlsx_error_text,
            "text_open_error": text_error_text,
        }

    return {
        "format": "unknown_binary",
        "bytes": len(payload),
        "magic_hex": payload[:16].hex(),
        "xlsx_open_error": xlsx_error_text,
        "text_open_error": text_error_text,
    }


def _coverage_candidates(schema: dict) -> list[dict]:
    if schema.get("format") == "xlsx":
        rows = []
        for sheet in schema.get("sheets", []):
            for year in sheet.get("year_columns", []):
                rows.append({"sheet": sheet.get("sheet"), **year})
        return rows
    return list(schema.get("year_columns", []))


def _covers(schema: dict, start: int, end: int) -> bool:
    for row in _coverage_candidates(schema):
        lo = row.get("min_year")
        hi = row.get("max_year")
        if lo is not None and hi is not None and lo <= start and hi >= end:
            return True
    return False


def _source_gate_status(
    contract: dict,
    migrant_schema: dict,
    resident_schema: dict,
    resource_schema: dict,
) -> tuple[str, list[str]]:
    reasons = []
    if not _covers(
        migrant_schema,
        contract["sources"]["migrant_timing"]["years"][0],
        contract["sources"]["migrant_timing"]["years"][1],
    ):
        reasons.append("MIGRANT_YEAR_COVERAGE_NOT_CERTIFIED")
    if not _covers(
        resident_schema,
        contract["sources"]["resident_partner_timing"]["years"][0],
        contract["sources"]["resident_partner_timing"]["years"][1],
    ):
        reasons.append("RESIDENT_YEAR_COVERAGE_NOT_CERTIFIED")
    if not _covers(
        resource_schema,
        contract["sources"]["destination_resource_state"]["years"][0],
        contract["sources"]["destination_resource_state"]["years"][1],
    ):
        reasons.append("RESOURCE_YEAR_COVERAGE_NOT_CERTIFIED")

    if reasons:
        return "SCHEMA_REVIEW_REQUIRED", reasons
    return "BIOLOGICAL_SOURCE_GATE_PASS_CUE_EXTENSION_PENDING", reasons


def _load_cue_extension_receipt(path: Path | None) -> tuple[dict, bool]:
    if path is None:
        return {"status": "NOT_PROVIDED"}, False
    if not path.exists():
        return {"status": "MISSING", "path": str(path)}, False

    payload = json.loads(path.read_text(encoding="utf-8"))
    firewall = payload.get("outcome_firewall", {})
    forbidden_true = [
        key
        for key, value in firewall.items()
        if value is True
    ]
    ok = (
        payload.get("status") == "SOURCE_FAITHFUL_CUE_EXTENSION_COMPLETE"
        and payload.get("years") == [1980, 2015]
        and payload.get("n_years") == 36
        and not forbidden_true
        and bool(payload.get("annual_csv_sha256"))
    )
    summary = {
        "status": payload.get("status"),
        "years": payload.get("years"),
        "n_years": payload.get("n_years"),
        "annual_csv_sha256": payload.get("annual_csv_sha256"),
        "cue_window": payload.get("cue_window"),
        "grid_latitudes_deg_n": payload.get("grid_latitudes_deg_n"),
        "grid_longitudes_deg_e": payload.get("grid_longitudes_deg_e"),
        "transport_counts": payload.get("transport_counts"),
        "outcome_firewall": firewall,
        "firewall_true_flags": forbidden_true,
        "certified": ok,
    }
    return summary, ok


def main():
    args = parse_args()
    contract = json.loads(args.contract.read_text(encoding="utf-8"))
    cue_extension, cue_extension_ok = _load_cue_extension_receipt(
        args.cue_receipt
    )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    sources_dir = args.output_dir / "source_files"
    if sources_dir.exists():
        shutil.rmtree(sources_dir)
    sources_dir.mkdir(parents=True)

    session = _session()

    receipt_path = args.output_dir / "source_gate_a_receipt.json"
    acquired = {}
    try:
        acquired["migrant_timing"] = _mda_download_file(
            session,
            contract["sources"]["migrant_timing"]["archive_url"],
            sources_dir,
            args.timeout,
        )
        acquired["migrant_timing"]["schema"] = inspect_source(
            Path(acquired["migrant_timing"]["path"])
        )

        acquired["resident_partner_timing"] = _dryad_download_file(
            session,
            contract["sources"]["resident_partner_timing"]["dataset_doi"],
            contract["sources"]["resident_partner_timing"]["file"],
            sources_dir,
            args.timeout,
        )
        acquired["resident_partner_timing"]["schema"] = inspect_source(
            Path(acquired["resident_partner_timing"]["path"])
        )

        acquired["destination_resource_state"] = _dryad_download_file(
            session,
            contract["sources"]["destination_resource_state"]["dataset_doi"],
            contract["sources"]["destination_resource_state"]["file"],
            sources_dir,
            args.timeout,
        )
        acquired["destination_resource_state"]["schema"] = inspect_source(
            Path(acquired["destination_resource_state"]["path"])
        )
    except Exception as exc:
        failure = {
            "result_id": "payoff_b_hoge_veluwe_source_gate_a_20260927",
            "contract_id": contract["contract_id"],
            "status": "SOURCE_ACCESS_FAILURE",
            "cue_extension": cue_extension,
            "error_type": type(exc).__name__,
            "error": str(exc),
            "acquired_sources_before_failure": {
                key: {
                    field: value
                    for field, value in meta.items()
                    if field != "path"
                }
                for key, meta in acquired.items()
            },
            "outcome_firewall": {
                "cross_source_join_performed": False,
                "focal_partner_mismatch_computed": False,
                "cue_resource_connectivity_computed": False,
                "information_reversal_gate_opened": False,
                "history_test_opened": False,
            },
        }
        receipt_path.write_text(
            json.dumps(failure, indent=2) + "\n",
            encoding="utf-8",
        )
        shutil.rmtree(sources_dir, ignore_errors=True)
        print(json.dumps(failure, indent=2))
        raise

    migrant = acquired["migrant_timing"]
    resident = acquired["resident_partner_timing"]
    resource = acquired["destination_resource_state"]

    migrant_schema = migrant["schema"]
    resident_schema = resident["schema"]
    resource_schema = resource["schema"]

    status, reasons = _source_gate_status(
        contract,
        migrant_schema,
        resident_schema,
        resource_schema,
    )
    if status == "BIOLOGICAL_SOURCE_GATE_PASS_CUE_EXTENSION_PENDING":
        if cue_extension_ok:
            status = "GATE_A_PASS_SOURCE_READY"
        else:
            status = "CUE_EXTENSION_NOT_SOURCE_FAITHFUL"
            reasons = [*reasons, "CUE_EXTENSION_NOT_CERTIFIED"]

    receipt = {
        "result_id": "payoff_b_hoge_veluwe_source_gate_a_20260927",
        "contract_id": contract["contract_id"],
        "status": status,
        "reasons": reasons,
        "registered_source_overlap": contract["population"]["primary_overlap_years"],
        "registered_history_span": contract["population"][
            "primary_history_years_after_connectivity_construction"
        ],
        "registered_history_year_count": contract["population"][
            "expected_primary_history_year_count"
        ],
        "site_alignment": {
            "registered_site": contract["population"]["site"],
            "resident_resource_basis": (
                "Dryad DOI 10.5061/dryad.f1vhhmgx6 is published for the "
                "Hoge Veluwe great-tit population and registered exact files"
            ),
            "migrant_basis": (
                "Tomotani et al. Marine Data Archive dataset linked to "
                "DOI 10.1111/gcb.14006 and the registered Hoge Veluwe lane"
            ),
            "joined_site_filter_applied": False,
        },
        "sources": {
            "migrant_timing": {
                "provenance": {
                    key: value
                    for key, value in migrant.items()
                    if key not in {"path", "schema"}
                },
                "schema": migrant_schema,
            },
            "resident_partner_timing": {
                "provenance": {
                    key: value
                    for key, value in resident.items()
                    if key not in {"path", "schema"}
                },
                "schema": resident_schema,
            },
            "destination_resource_state": {
                "provenance": {
                    key: value
                    for key, value in resource.items()
                    if key not in {"path", "schema"}
                },
                "schema": resource_schema,
            },
        },
        "cue_extension": {
            **cue_extension,
            "registered_rule": contract["sources"]["precommitment_cue"]["spatial_rule"],
        },
        "outcome_firewall": {
            "cross_source_join_performed": False,
            "focal_partner_mismatch_computed": False,
            "cue_resource_connectivity_computed": False,
            "information_reversal_gate_opened": False,
            "history_test_opened": False,
            "migrant_timing_values_exported": False,
            "resident_timing_values_exported": False,
            "resource_peak_values_exported": False,
        },
        "claim_boundary": (
            "Gate A certifies source identity, immutable bytes, schema and year "
            "coverage only. It is not an empirical hysteresis result."
        ),
    }

    # Do not retain raw source files in the uploaded Gate-A artifact.
    receipt_path.write_text(
        json.dumps(receipt, indent=2) + "\n",
        encoding="utf-8",
    )
    shutil.rmtree(sources_dir)

    print(
        "PAYOFF_B_HV_SOURCE_GATE_A "
        f"status={status} "
        f"migrant_sha256={migrant['sha256']} "
        f"resident_sha256={resident['sha256']} "
        f"resource_sha256={resource['sha256']}"
    )
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    main()
