#!/usr/bin/env python3
"""Schema-only probe for Wang et al. 2024 Figshare archive.

This script deliberately does NOT report numerical timing values. It records
public file metadata, sheet/table headers and row counts so the precommitted
eligibility gate can be evaluated before outcome analysis.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path

import requests

ARTICLE_ID = 24613599
API = f"https://api.figshare.com/v2/articles/{ARTICLE_ID}"
OUT = Path("outputs/payoff_b_wang_figshare_schema_probe.json")

def norm(x):
    return re.sub(r"[^a-z0-9]+", "_", str(x).strip().lower()).strip("_")

def inspect_csv(raw: bytes):
    text = raw.decode("utf-8-sig", errors="replace")
    rows = csv.reader(io.StringIO(text))
    header = next(rows, [])
    n = sum(1 for _ in rows)
    return {"kind":"csv","headers":[norm(x) for x in header],"data_rows":n}

def inspect_xlsx(raw: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    out=[]
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1,max_row=1,values_only=True), ())
        out.append({
            "sheet":ws.title,
            "headers":[norm(x) for x in first],
            "max_row":ws.max_row,
            "max_column":ws.max_column,
        })
    return {"kind":"xlsx","sheets":out}

def inspect_zip(raw: bytes):
    z=zipfile.ZipFile(io.BytesIO(raw))
    return {"kind":"zip","members":[{"name":i.filename,"size":i.file_size} for i in z.infolist()]}

def main():
    meta = requests.get(API, timeout=60)
    meta.raise_for_status()
    js = meta.json()
    result={
        "article_id":ARTICLE_ID,
        "title":js.get("title"),
        "doi":js.get("doi"),
        "license":js.get("license"),
        "files":[],
        "policy":"SCHEMA_ONLY_NO_NUMERICAL_TIMING_VALUES",
    }
    for f in js.get("files",[]):
        rec={k:f.get(k) for k in ("id","name","size","download_url","supplied_md5")}
        name=(f.get("name") or "").lower()
        if name.endswith((".csv",".tsv",".xlsx",".zip")) and f.get("download_url"):
            r=requests.get(f["download_url"],timeout=120)
            r.raise_for_status()
            raw=r.content
            if name.endswith(".csv"):
                rec["schema"]=inspect_csv(raw)
            elif name.endswith(".tsv"):
                rec["schema"]=inspect_csv(raw.replace(b"\t",b","))
            elif name.endswith(".xlsx"):
                rec["schema"]=inspect_xlsx(raw)
            elif name.endswith(".zip"):
                rec["schema"]=inspect_zip(raw)
        result["files"].append(rec)
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding="utf-8")
    print(json.dumps(result,indent=2,ensure_ascii=False))

if __name__=="__main__":
    main()
