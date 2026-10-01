from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def _set_font(run, size: Pt = Pt(11), bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = size
    run.font.color.rgb = None
    run.font.color.theme_color = MSO_THEME_COLOR_INDEX.TEXT_1
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Times New Roman")


def _suppress_line_numbers(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:suppressLineNumbers")) is None:
        ppr.append(OxmlElement("w:suppressLineNumbers"))


def _set_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    _suppress_line_numbers(paragraph)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, cached, end])
    _set_font(run, Pt(9))


def _set_line_numbers(section, enabled: bool) -> None:
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:lnNumType"))
    if old is not None:
        sect_pr.remove(old)
    if not enabled:
        return
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:start"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sect_pr.append(ln)


def _style_document(
    doc: Document,
    *,
    line_numbers: bool,
    page_numbers: bool,
    compact: bool,
) -> None:
    line_spacing = 1.15 if compact else 1.5
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = line_spacing
    normal.paragraph_format.space_after = Pt(0)

    style_sizes = {
        "Title": Pt(14),
        "Heading 1": Pt(12),
        "Heading 2": Pt(11),
        "Heading 3": Pt(11),
        "Heading 4": Pt(11),
    }
    for style_name, size in style_sizes.items():
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = size
        style.font.bold = True
        style.font.color.rgb = None
        style.font.color.theme_color = MSO_THEME_COLOR_INDEX.TEXT_1
        style.paragraph_format.line_spacing = line_spacing
        style.paragraph_format.space_after = Pt(0)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing = line_spacing
        paragraph.paragraph_format.space_after = Pt(0)
        text_value = paragraph.text.strip()
        if line_numbers and text_value == "Figure legends":
            paragraph.paragraph_format.page_break_before = True
        if line_numbers and text_value.startswith("Figure "):
            match = re.match(r"Figure\s+(\d+)\.", text_value)
            if match:
                number = int(match.group(1))
                if number >= 2:
                    paragraph.paragraph_format.page_break_before = True
                paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            _set_font(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        _set_font(run, Pt(10))

    max_width = Inches(6.2)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        _set_line_numbers(section, line_numbers)
        if page_numbers:
            section.footer.is_linked_to_previous = False
            footer = section.footer
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            _set_page_field(p)

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = ""
    props.subject = ""
    props.keywords = ""
    props.comments = ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--line-numbers", action="store_true")
    parser.add_argument("--no-page-numbers", action="store_true")
    parser.add_argument("--compact", action="store_true")
    args = parser.parse_args()

    doc = Document(args.input_docx)
    _style_document(
        doc,
        line_numbers=args.line_numbers,
        page_numbers=not args.no_page_numbers,
        compact=args.compact,
    )
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output_docx)


if __name__ == "__main__":
    main()
