from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def _set_font(run, size: Pt = Pt(11)) -> None:
    run.font.name = "Times New Roman"
    run.font.size = size
    run.font.color.rgb = None
    run.font.color.theme_color = MSO_THEME_COLOR_INDEX.TEXT_1
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Times New Roman")


def _suppress_line_numbers(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:suppressLineNumbers")) is None:
        ppr.append(OxmlElement("w:suppressLineNumbers"))


def _set_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    _suppress_line_numbers(paragraph)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    _set_font(run, Pt(9))


def _set_line_numbers(section) -> None:
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:lnNumType"))
    if old is not None:
        sect_pr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:start"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sect_pr.append(ln)


def format_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(11)
            style.font.color.rgb = None
            style.font.color.theme_color = MSO_THEME_COLOR_INDEX.TEXT_1
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(0)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            _set_font(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.15
                    for run in paragraph.runs:
                        _set_font(run, Pt(10))

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        _set_line_numbers(section)
        section.footer.is_linked_to_previous = False
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        _set_page_field(p)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()
    doc = Document(args.input_docx)
    format_document(doc)
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output_docx)


if __name__ == "__main__":
    main()
