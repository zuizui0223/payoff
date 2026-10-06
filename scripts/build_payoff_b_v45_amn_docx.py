#!/usr/bin/env python3
"""Build an anonymous American Naturalist review DOCX for the V4.5 manuscript."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript" / "PAYOFF_B_FORECAST_ACCESS_CORRECTION_V4_5_AMNAT.md"
FRONT = ROOT / "submission" / "AMNAT_V4_5_FRONTMATTER_20261006.md"
CAPTIONS = ROOT / "submission" / "AMNAT_V4_5_FIGURE_CAPTIONS_20261006.md"

FORBIDDEN = (
    "PAYOFF-B", "V2", "V3", "V4", "V5", "V8",
    "post-freeze", "rollback", "workflow ID", "artifact ID",
)


def between(text: str, start: str, end: str) -> str:
    a = text.index(start) + len(start)
    b = text.index(end, a)
    return text[a:b].strip()


def clean(text: str) -> str:
    return text.replace("**", "").replace(chr(96), "")


def set_double(paragraph) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def add_page_number(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_line_numbers(section) -> None:
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:lnNumType"))
    if old is not None:
        sect_pr.remove(old)
    el = OxmlElement("w:lnNumType")
    el.set(qn("w:countBy"), "1")
    el.set(qn("w:start"), "1")
    el.set(qn("w:restart"), "continuous")
    sect_pr.append(el)


def add_heading(doc, value: str, level: int) -> None:
    p = doc.add_paragraph()
    set_double(p)
    r = p.add_run(clean(value))
    r.bold = True
    r.font.size = Pt(12 if level > 1 else 13)


def add_body(doc, value: str, indent: bool = True) -> None:
    p = doc.add_paragraph()
    set_double(p)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    p.add_run(clean(value))


def add_equation(doc, lines: list[str]) -> None:
    p = doc.add_paragraph()
    set_double(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(" ".join(x.strip() for x in lines if x.strip()))


def render_markdown(doc, text: str) -> None:
    lines = text.splitlines()
    buf: list[str] = []
    i = 0

    def flush() -> None:
        nonlocal buf
        if buf:
            add_body(doc, " ".join(x.strip() for x in buf))
            buf = []

    while i < len(lines):
        raw = lines[i]
        x = raw.strip()

        if not x or x == "---":
            flush()
            i += 1
            continue

        if x == "[":
            flush()
            eq = []
            i += 1
            while i < len(lines) and lines[i].strip() != "]":
                eq.append(lines[i])
                i += 1
            add_equation(doc, eq)
            i += 1
            continue

        if x.startswith("## "):
            flush()
            add_heading(doc, x[3:], 1)
            i += 1
            continue

        if x.startswith("### "):
            flush()
            add_heading(doc, x[4:], 2)
            i += 1
            continue

        if x.startswith("> "):
            flush()
            p = doc.add_paragraph()
            set_double(p)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            r = p.add_run(clean(x[2:]))
            r.italic = True
            i += 1
            continue

        if re.match(r"^\d+\.\s+", x):
            flush()
            add_body(doc, x, indent=False)
            i += 1
            continue

        buf.append(raw)
        i += 1

    flush()


def build(output: Path) -> dict:
    manuscript = MANUSCRIPT.read_text(encoding="utf-8")
    front = FRONT.read_text(encoding="utf-8")
    captions = CAPTIONS.read_text(encoding="utf-8")

    bad = [x for x in FORBIDDEN if x.lower() in manuscript.lower()]
    if bad:
        raise ValueError(f"anonymous manuscript contains internal labels: {bad}")

    title = manuscript.splitlines()[0].lstrip("# ").strip()
    short_title = clean(between(front, "## Short title", "## Abstract"))
    abstract = between(manuscript, "## Abstract", "Keywords:")
    keyword_block = manuscript.split("Keywords:", 1)[1].split("---", 1)[0]
    keywords = clean(keyword_block).replace("\n", " ").strip()
    keyword_list = [x.strip() for x in keywords.split(";") if x.strip()]

    if len(short_title) > 40:
        raise ValueError("short title exceeds 40 characters")
    if not (1 <= len(keyword_list) <= 6):
        raise ValueError("keyword count must be 1 to 6")

    main_text = manuscript[manuscript.index("## 1. Introduction"):]

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    add_line_numbers(sec)
    add_page_number(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    set_double(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)

    add_body(doc, f"Short title: {short_title}", indent=False)
    add_body(doc, f"Keywords: {keywords}", indent=False)
    add_body(doc, "Article type: Major Article", indent=False)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "Abstract", 1)
    add_body(doc, abstract.replace("\n", " "), indent=False)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    render_markdown(doc, main_text)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "Figure Legends", 1)
    cap_body = "\n".join(
        line for line in captions.splitlines()
        if not line.startswith("# ")
    )
    render_markdown(doc, cap_body)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

    return {
        "output": str(output),
        "title": title,
        "short_title": short_title,
        "keywords": len(keyword_list),
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--output",
        type=Path,
        default=Path("outputs/payoff_b_v45_amn_anonymous_manuscript.docx"),
    )
    args = ap.parse_args()
    result = build(args.output)
    print(result)


if __name__ == "__main__":
    main()
