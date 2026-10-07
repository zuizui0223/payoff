import importlib.util
import sys
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "build_payoff_b_v45_amn_docx.py"


def load_module():
    spec = importlib.util.spec_from_file_location("v45docx", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def test_docx_build_and_structure(tmp_path):
    module = load_module()
    out = tmp_path / "anonymous.docx"
    result = module.build(out)
    assert out.exists()
    assert result["keywords"] <= 6
    assert len(result["short_title"]) <= 40
    assert 1000 <= result["text_words"] <= 7500

    doc = Document(out)
    full = "\n".join(p.text for p in doc.paragraphs)

    assert "Seasonal tracking depends on information access" in full
    assert "Short title: Information access and correction" in full
    assert "Word count excluding Literature Cited:" in full
    assert "Manuscript elements:" in full
    assert "Abstract" in full
    assert "1. Introduction" in full
    assert "3. Methods" in full
    assert "Figure Legends" in full
    assert "Figure 1." in full
    assert "Figure 2." in full
    assert "Figure 3." in full

    for forbidden in module.FORBIDDEN:
        assert forbidden.lower() not in full.lower()

    # raw LaTeX control sequences should not leak into the review document
    assert "\\Delta" not in full
    assert "\\text{" not in full
    assert "\\!" not in full
    assert "Delta" not in full
    assert "e_{t+1}" not in full
    assert "eₜ₊₁" in full
    assert "Δ e_route" in full
    assert "σ_Y²" in full
    assert "ρ²" in full
    assert "q(t)=q₀+Δ q[1-exp(-α t)]" in full


def test_docx_has_line_and_page_number_fields(tmp_path):
    module = load_module()
    out = tmp_path / "anonymous.docx"
    module.build(out)

    with zipfile.ZipFile(out) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
        footer_names = [n for n in zf.namelist() if n.startswith("word/footer")]
        footer_xml = "".join(zf.read(n).decode("utf-8") for n in footer_names)

    assert "w:lnNumType" in document_xml
    assert 'w:countBy="1"' in document_xml
    assert " PAGE " in footer_xml


def test_docx_normal_style_is_double_spaced(tmp_path):
    module = load_module()
    out = tmp_path / "anonymous.docx"
    module.build(out)
    doc = Document(out)
    assert doc.styles["Normal"].paragraph_format.line_spacing == 2.0
